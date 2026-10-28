#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
tests/test_fuse_types.py
========================
Fuse-type resolution: the rated current comes from the source data
(``config["fuse"]["current_field"]``, default ``T_jar_nami``) and selects the
matching ``TypFuse`` entry of the bundled Excel type library.

Most cases exercise :meth:`Converter.fuse_type_name` / ``fuse_type_id`` on a
Converter instance built without ``__init__`` (no shapefiles, no license), so
the resolution rules are tested in isolation; the last test runs a complete
conversion and checks the emitted ``RelFuse``/``TypFuse`` rows.
"""
import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from conftest_free_harness import unlocked, write_layer, temp_dir

from shp2dgs.core.config_io import (
    canonicalize_type_sheet_columns, load_default_types_override, load_json,
    merge_type_library,
)
from shp2dgs.core.constants import TYPE_LIBRARY_SHEETS
from shp2dgs.core.converter import Converter

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def col(header, name):
    """Index of a DGS column by name (headers look like ``loc_name(a:40)``)."""
    return next(i for i, h in enumerate(header) if str(h).split("(")[0] == name)


def _bare_converter(cfg=None, tpl=None):
    """A Converter with just the state the fuse logic needs (no __init__)."""
    cfg = cfg or load_json(os.path.join(REPO_ROOT, "config.json"))
    tpl = tpl or load_json(os.path.join(REPO_ROOT, "dgs_template.json"))
    override = load_default_types_override(tpl)
    merged = merge_type_library(tpl, override)
    for sheet in TYPE_LIBRARY_SHEETS:
        if sheet in merged:
            merged[sheet] = canonicalize_type_sheet_columns(merged[sheet])

    conv = Converter.__new__(Converter)
    conv.cfg = cfg
    conv.tpl = merged
    conv.uknom = float(cfg["nominal_mv_kv"])
    conv.lv_kv = float(cfg["transformer"]["lv_voltage_kv"])
    conv.typfuse_id = {}
    header = merged["TypFuse"]["header"]
    id_idx = col(header, "ID")
    name_idx = col(header, "loc_name")
    for row in merged["TypFuse"]["rows"]:
        conv.typfuse_id[row[name_idx]] = row[id_idx]
    conv.issues = []
    conv.manual_overrides = {}
    conv._lv_fuse_type_rows = []
    conv._lv_fuse_type_ids = {}
    conv._next_id = 9_000_000
    conv._termcount = 0
    return conv


class FuseNameResolutionTests(unittest.TestCase):
    def setUp(self):
        self.conv = _bare_converter()

    def test_exact_rating_from_data(self):
        for amps in (1, 8, 63, 100):
            name = self.conv.fuse_type_name({"T_jar_nami": amps})
            self.assertEqual(name, f"Fuse_Type_{amps}")
            self.assertEqual(self.conv.issues, [], "an exact match must not report an issue")

    def test_rating_as_text_decorated_and_persian_digits(self):
        for raw, expected in (("63", 63), ("63 A", 63), ("63.0", 63),
                              ("فیوز ۶۳ آمپر", 63), ("۲۵", 25)):
            name = self.conv.fuse_type_name({"T_jar_nami": raw})
            self.assertEqual(name, f"Fuse_Type_{expected}", f"raw={raw!r}")

    def test_unlisted_rating_snaps_down_and_is_reported(self):
        # 55 A is not in the library (…50, 63…); the lower rating is used,
        # because a fuse rated above the device would under-protect it.
        name = self.conv.fuse_type_name({"T_jar_nami": 55})
        self.assertEqual(name, "Fuse_Type_50")
        self.assertEqual([i["category"] for i in self.conv.issues], ["fuse_rating_snapped"])

    def test_snap_prefer_nearest(self):
        cfg = load_json(os.path.join(REPO_ROOT, "config.json"))
        cfg["fuse"]["snap_prefer"] = "nearest"
        conv = _bare_converter(cfg=cfg)
        self.assertEqual(conv.fuse_type_name({"T_jar_nami": 55}), "Fuse_Type_50")
        self.assertEqual(conv.fuse_type_name({"T_jar_nami": 58}), "Fuse_Type_63")

    def test_missing_or_unusable_current_uses_configured_fallback(self):
        for rec in ({}, {"T_jar_nami": None}, {"T_jar_nami": ""}, {"T_jar_nami": "n/a"}):
            conv = _bare_converter()
            name = conv.fuse_type_name(rec)
            self.assertEqual(name, self.conv.cfg["fuse"]["type_name"])
            self.assertEqual([i["category"] for i in conv.issues], ["fuse_current_missing"])

    def test_custom_current_field_name(self):
        cfg = load_json(os.path.join(REPO_ROOT, "config.json"))
        cfg["fuse"]["current_field"] = "T_amper"
        conv = _bare_converter(cfg=cfg)
        self.assertEqual(conv.fuse_type_name({"T_amper": 24, "T_jar_nami": 1}),
                         "Fuse_Type_24")

    def test_lv_fuse_is_synthesized_at_its_own_rating(self):
        """The shipped library is MV-only (urat 20 kV), so LV fuses must not
        borrow an MV type -- one LV type per rating is created instead."""
        tid = self.conv.fuse_type_id({"T_jar_nami": 25}, "lv")
        rows = {r[1]: r for r in self.conv._lv_fuse_type_rows}
        self.assertIn("LV_Fuse_25A", rows)
        self.assertEqual(rows["LV_Fuse_25A"][3], 0.4)   # urat
        self.assertEqual(rows["LV_Fuse_25A"][4], 25)    # irat
        self.assertEqual(tid, self.conv.typfuse_id["LV_Fuse_25A"])
        # reused, not duplicated, on the next occurrence
        self.conv.fuse_type_id({"T_jar_nami": 25}, "lv")
        self.assertEqual(len(self.conv._lv_fuse_type_rows), 1)

    def test_lv_fuse_without_data_uses_config_default(self):
        self.conv.fuse_type_id({}, "lv")
        self.assertEqual([r[1] for r in self.conv._lv_fuse_type_rows], ["LV_Default_Fuse"])


class EndToEndFuseTests(unittest.TestCase):
    def test_conversion_emits_rated_fuse_types(self):
        import openpyxl

        shp_dir = temp_dir()
        write_layer(shp_dir, "ELWFWT",
                    [("T_MapID", "C", 40, 0), ("Feeder", "C", 40, 0)],
                    [{"T_MapID": "SRC1", "Feeder": "F1 - North"}],
                    [(0.0, 0.0)])
        write_layer(shp_dir, "ELWOML",
                    [("T_MapID", "C", 40, 0), ("Feeder", "C", 40, 0), ("T_jns_hadi", "C", 20, 0),
                     ("T_cod_sath", "N", 10, 0), ("T_typ_shab", "C", 20, 0), ("Shape_STLe", "N", 12, 2)],
                    [{"T_MapID": "L1", "Feeder": "F1 - North", "T_jns_hadi": "ACSR",
                      "T_cod_sath": 70, "T_typ_shab": "", "Shape_STLe": 1000.0}],
                    [[(0.0, 0.0), (1000.0, 0.0)]])
        write_layer(shp_dir, "ELWFCT",
                    [("T_MapID", "C", 40, 0), ("Feeder", "C", 40, 0), ("T_jar_nami", "N", 10, 0)],
                    [{"T_MapID": "FU63", "Feeder": "F1 - North", "T_jar_nami": 63}],
                    [(500.0, 0.0)])
        write_layer(shp_dir, "ELWDTR",
                    [("T_MapID", "C", 40, 0), ("Feeder", "C", 40, 0), ("T_strn", "N", 10, 0),
                     ("T_tozieh", "C", 40, 0), ("T_noe_tera", "C", 40, 0)],
                    [{"T_MapID": "TR1", "Feeder": "F1 - North", "T_strn": 100,
                      "T_tozieh": "کم تلفات", "T_noe_tera": "روغن"}],
                    [(1000.0, 0.0)])
        write_layer(shp_dir, "ELWOLL",
                    [("T_MapID", "C", 40, 0), ("Feeder", "C", 40, 0), ("T_jns_hadi", "C", 20, 0),
                     ("T_cod_sath", "N", 10, 0), ("T_typ_shab", "C", 20, 0), ("Shape_STLe", "N", 12, 2)],
                    [{"T_MapID": "LVL1", "Feeder": "F1 - North", "T_jns_hadi": "مس",
                      "T_cod_sath": 35, "T_typ_shab": "لخت", "Shape_STLe": 50.0}],
                    [[(1000.0, 0.0), (1050.0, 0.0)]])
        write_layer(shp_dir, "ELWLFU",
                    [("T_MapID", "C", 40, 0), ("Feeder", "C", 40, 0), ("T_jar_nami", "N", 10, 0)],
                    [{"T_MapID": "LVFU25", "Feeder": "F1 - North", "T_jar_nami": 25}],
                    [(1025.0, 0.0)])

        out_path = os.path.join(shp_dir, "out.xlsx")
        cfg = load_json(os.path.join(REPO_ROOT, "config.json"))
        tpl = load_json(os.path.join(REPO_ROOT, "dgs_template.json"))
        with unlocked():
            conv = Converter(shp_dir, cfg, tpl, make_graphics=False,
                             type_library_override=load_default_types_override(tpl))
            conv.build().emit().write(out_path)

        wb = openpyxl.load_workbook(out_path, data_only=True)
        fuse_rows = list(wb["TypFuse"].iter_rows(values_only=True))
        fuse_hdr = fuse_rows[0]
        by_name = {r[col(fuse_hdr, "loc_name")]: r for r in fuse_rows[1:]}

        rel_rows = list(wb["RelFuse"].iter_rows(values_only=True))
        rel_hdr = rel_rows[0]
        used_types = {r[col(rel_hdr, "typ_id")] for r in rel_rows[1:]}
        mv_fuse = by_name["Fuse_Type_63"]
        lv_fuse = by_name["LV_Fuse_25A"]
        self.assertIn(mv_fuse[col(fuse_hdr, "ID")], used_types)
        self.assertIn(lv_fuse[col(fuse_hdr, "ID")], used_types)
        self.assertEqual(mv_fuse[col(fuse_hdr, "urat")], 20)
        self.assertEqual(lv_fuse[col(fuse_hdr, "urat")], 0.4)
        self.assertEqual(lv_fuse[col(fuse_hdr, "irat")], 25)

        # The Excel library's values (not dgs_template's placeholders) must
        # reach the deck: 100 kVA low-loss transformer, 70 mm² ACSR line.
        tr_rows = list(wb["TypTr2"].iter_rows(values_only=True))
        tr_hdr = tr_rows[0]
        tr = next(r for r in tr_rows[1:]
                  if r[col(tr_hdr, "loc_name")] == "TRF-AB'-20/0.4 kV-100 kVA")
        self.assertAlmostEqual(tr[col(tr_hdr, "pcutr")], 1.75)
        lne_rows = list(wb["TypLne"].iter_rows(values_only=True))
        lne_hdr = lne_rows[0]
        mink = next(r for r in lne_rows[1:] if r[col(lne_hdr, "loc_name")] == "ACSR-MINK")
        self.assertAlmostEqual(mink[col(lne_hdr, "rline")], 0.46)


if __name__ == "__main__":
    unittest.main()
